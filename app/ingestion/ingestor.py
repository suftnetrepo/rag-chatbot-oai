"""
Document ingestion pipeline.

Steps per document:
  1. Parse raw bytes into text (PDF, DOCX, TXT, MD, CSV).
  2. Split text into overlapping chunks via LlamaIndex SentenceSplitter.
  3. Embed each chunk with OpenAI text-embedding-3-large.
  4. Upsert chunks into Typesense (vector + BM25 fields).
  5. Return chunk data for the caller to persist to PostgreSQL.
"""
from __future__ import annotations

import io
import uuid
from datetime import datetime, timezone
from pathlib import Path

from llama_index.core.node_parser import SentenceSplitter
from llama_index.core.schema import Document as LIDocument

from app.agents.llm_factory import get_embed_model
from app.config import get_settings
from app.logging_config import get_logger
from app.retrieval.typesense_client import get_typesense_client, upsert_chunks

log = get_logger(__name__)

ALLOWED_EXTENSIONS = {".pdf", ".txt", ".md", ".csv", ".docx"}


def _extract_text(filename: str, content: bytes) -> str:
    ext = Path(filename).suffix.lower()
    if ext in (".txt", ".md", ".csv"):
        return content.decode("utf-8", errors="replace")
    if ext == ".pdf":
        try:
            import pypdf
            reader = pypdf.PdfReader(io.BytesIO(content))
            return "\n\n".join(p.extract_text() or "" for p in reader.pages)
        except ImportError:
            log.warning("pypdf_not_installed", filename=filename)
            return content.decode("utf-8", errors="replace")
    if ext == ".docx":
        try:
            import docx
            doc = docx.Document(io.BytesIO(content))
            return "\n\n".join(p.text for p in doc.paragraphs if p.text)
        except ImportError:
            log.warning("python_docx_not_installed", filename=filename)
            return content.decode("utf-8", errors="replace")
    raise ValueError(f"Unsupported file type: {ext}")


def _chunk_text(text: str, filename: str, document_id: str):
    settings = get_settings()
    splitter = SentenceSplitter(
        chunk_size=settings.chunk_size,
        chunk_overlap=settings.chunk_overlap,
    )
    base_doc = LIDocument(
        text=text,
        metadata={"filename": filename, "document_id": document_id},
    )
    return splitter.get_nodes_from_documents([base_doc])


class DocumentIngestor:
    """
    Stateless ingestor — receives raw bytes, returns a result dict.
    No ORM objects are touched here — pure data in, pure data out.
    """

    def __init__(self) -> None:
        self._embed_model = get_embed_model()
        self._ts_client   = get_typesense_client()

    def ingest(self, filename: str, content: bytes, document_id: str) -> dict:
        """
        Args:
            filename:    Original filename.
            content:     Raw file bytes.
            document_id: UUID string of the pre-created Document record.

        Returns:
            {
                "status":       "indexed" | "failed",
                "chunk_count":  int,
                "error":        str | None,
                "chunks":       list[dict],   # for DB persistence by caller
                "indexed_at":   datetime,
            }
        """
        log.info("ingestion_start", filename=filename, doc_id=document_id)

        try:
            # 1. Extract text
            text = _extract_text(filename, content)
            if not text.strip():
                raise ValueError("Extracted text is empty — document may be image-only.")

            # 2. Chunk
            nodes = _chunk_text(text, filename, document_id)
            log.info("chunks_created", count=len(nodes), doc_id=document_id)

            # 3. Embed
            texts_to_embed = [node.get_content() for node in nodes]
            embeddings     = self._embed_model.get_text_embedding_batch(
                texts_to_embed, show_progress=False
            )

            # 4. Build Typesense + DB records
            ts_docs   = []
            db_chunks = []

            for idx, (node, embedding) in enumerate(zip(nodes, embeddings)):
                chunk_text  = node.get_content()
                ts_id       = f"{document_id}__{idx}"
                token_count = len(chunk_text.split())

                ts_docs.append({
                    "id":          ts_id,
                    "document_id": document_id,
                    "filename":    filename,
                    "chunk_index": idx,
                    "content":     chunk_text,
                    "token_count": token_count,
                    "embedding":   embedding,
                })
                db_chunks.append({
                    "typesense_id":   ts_id,
                    "chunk_index":    idx,
                    "content":        chunk_text,
                    "token_count":    token_count,
                    "chunk_metadata": node.metadata,
                })

            # 5. Upsert into Typesense
            upsert_chunks(self._ts_client, ts_docs)

            log.info("ingestion_complete", doc_id=document_id, chunks=len(db_chunks))
            return {
                "status":      "indexed",
                "chunk_count": len(db_chunks),
                "error":       None,
                "chunks":      db_chunks,
                "indexed_at":  datetime.now(timezone.utc),
            }

        except Exception as exc:
            log.error("ingestion_failed", doc_id=document_id, error=str(exc))
            return {
                "status":      "failed",
                "chunk_count": 0,
                "error":       str(exc),
                "chunks":      [],
                "indexed_at":  None,
            }


_ingestor: DocumentIngestor | None = None


def get_ingestor() -> DocumentIngestor:
    global _ingestor
    if _ingestor is None:
        _ingestor = DocumentIngestor()
    return _ingestor
